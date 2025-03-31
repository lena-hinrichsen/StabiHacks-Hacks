import json
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font
import xml.etree.ElementTree as ET
from deep_translator import GoogleTranslator
from googletrans import Translator
from nltk.corpus import wordnet
import wn
import glob

object_dir = 'image-analysis\\results\\yolo9000\\'
text_dir = 'text_analysis\\results'
fulltext_dir = 'sbbget\\sbbget_downloads\download_temp\\'
ppn_list = 'text_image\\test-ten-books.txt'


def all_objects_in_corpus():
    objects_in_corpus = []
    for subdir, dirs, files in os.walk(object_dir):
        for filename in files:
            ppn = filename[0:12]
            objects = get_objects(ppn)
            for x in range(0,len(objects)):
                if objects[x] not in objects_in_corpus:
                    objects_in_corpus.append(objects[x])
                    objects_in_corpus.sort()
                    print(objects_in_corpus)
                    return(all_objects_in_corpus)


def get_objects(ppn):
    all_objects = []
    file = object_dir + ppn + '_objects.json'
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        y = json.loads(data)
        for x in range(0, len(y)):
            element = y[x]
            z = element.get("objects")
            for a in range(0, len(z)):
                b = z[a]
                obj = b.get('name')
                if obj not in all_objects:
                    all_objects.append(obj)
    print(all_objects)
    return all_objects


def get_class_id_from_names(object_name):
    file = open('text_image\\9k.names')
    list = file.read().splitlines()
    class_id = None
    for x in range(0, len(list)):
        if list[x] == object_name:
            class_id = x
            return class_id


def get_class_id_from_label(label):
    file = open('text_image\\9k.labels')
    list = file.read().splitlines()
    class_id = None
    for x in range(0, len(list)):
        if list[x] == label:
            class_id = x
            return class_id


def get_synset_id(object_name):
    class_id = get_class_id_from_names(object_name)
    synset_id = None
    if class_id == None:
        "Error! Maybe you typed in a non existent object_class. Please look up the list of possible objects in " \
        "9k.names!"
    else:
        file = open('text_image\\9k.labels')
        list = file.read().splitlines()
        synset_id = list[class_id]
        print(synset_id)
    return synset_id


def get_hyponyms(word, medium):
    wnetIDs = []
    all_lemmas = []
    if medium == 'image':
        id = get_synset_id(word)
        a = re.search(r"\d+(\.\d+)?", id)
        id_str = a.group(0)
        id_int = int(id_str)
        s = wordnet.synset_from_pos_and_offset('n', id_int)
        hypos = lambda s: s.hyponyms()
        hypos_list = list(s.closure(hypos))
        #print(hypos_list)
        all_lemmas = []
        wnetIDs = []
        all_lemmas_names = []
        for x in range(0, len(hypos_list)):
            # for images, we are interested in the class IDs of the hyponyms
            ss = hypos_list[x]
            current_offset = ss.offset()
            current_pos = ss.pos()
            offset_8 = str(current_offset).zfill(8)
            wnetID = current_pos + offset_8
            wnetIDs.append(wnetID)
    # for text, we are interested in the lemmas
    if medium == 'text':
        all_synsets = wn.synsets(word, lang='de')
        for x in range(0, len(all_synsets)):
            # currently, you cant search for hyponyms of hyponyms ... for future dev:
            # hypos = get_all_hyponyms_from_synset_odenet(all_synsets[x])
            hypos = all_synsets[x].hyponyms()
            for y in range(0, len(hypos)):
                lemmas = hypos[y].lemmas()
                for z in range(0, len(lemmas)):
                    if lemmas[z] not in all_lemmas:
                        all_lemmas.append(lemmas[z])
    if medium == 'image':
        return(wnetIDs)
    if medium == 'text':
        return(all_lemmas)


def get_all_hyponyms_from_id_wordnet(id):
    s = wordnet.synset_from_pos_and_offset('n', id)
    hypos = lambda s: s.hyponyms()
    hypos_list = list(s.closure(hypos))
    return hypos_list

# not ready yet
def get_all_hyponyms_from_synset_odenet(synset):
    hypos = lambda s: s.hyponyms()
    hypos_list = list(synset.closure(hypos))
    print(hypos_list)
    return hypos_list

#print(get_hyponyms('animal', 'image'))


def get_pages_with_object_illu(object_name, ppn, hyponyms):
    all_pages = []
    file = object_dir + ppn + '_objects.json'
    id = get_class_id_from_names(object_name)
    # print(file)
    obj_class_list = get_hyponyms(object_name, 'image')
    used_words_for_obj = []
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        y = json.loads(data)
        for x in range(0, len(y)):
            element = y[x]
            z = element.get("objects")
            for a in range(0, len(z)):
                b = z[a]
                current_id = b.get('class_id')
                current_obj = b.get('name')
                print(current_id)
                if current_id == id:
                    page = y[x].get('filename')
                    if current_obj not in used_words_for_obj:
                        used_words_for_obj.append(current_obj)
                    if page not in all_pages:
                        all_pages.append(page)
                if hyponyms:
                    print(obj_class_list)
                    print(y[x])
                    print(z[a])
                    current_id_wnet = get_synset_id(current_obj)
                    print("Current_id_wnet")
                    print(current_id_wnet)
                    if current_id_wnet in obj_class_list:
                        if current_obj not in used_words_for_obj:
                            used_words_for_obj.append(current_obj)
                        #print(y[x])
                        page = y[x].get('filename')
                        print(page)
                        #print(page)
                        if page not in all_pages:
                            all_pages.append(page)
    print(all_pages, used_words_for_obj)
    return(all_pages, used_words_for_obj)


def get_pages_with_named_entity(object_name, ppn):
    all_pages = []
    all_objects = []
    current_page = "1"
    for x in range(1, 9999):
        current_page = str(x)
        current_page = current_page.zfill(4)
        dir = fulltext_dir + ppn + '\\' + 'FILE_' + current_page + '_FULLTEXT' + '\\'
        # print(dir)
        fulltext_file_path = dir + current_page.zfill(8) + '.txt'
        json_file_path = dir + current_page.zfill(8) + '_ner_details.json'
        if os.path.exists(dir):
            if os.path.exists(json_file_path):
                with open(json_file_path) as f:
                    data = json.load(f)
                    for y in range(0, len(data['entities'])):
                        current_element = data['entities'][y]
                        # print(current_element)
                        current_text = current_element['text']
                        if len(current_element['labels']) == 1:
                            current_value = current_element['labels'][0]
                            # print(current_value)
                            if current_value['_value'] == object_name:
                                if current_page not in all_pages:
                                    all_pages.append(current_page)
                                if current_text not in all_objects:
                                    all_objects.append(current_text)
                        else:
                            print('There seems to be more than one label')
        else:
            # print('doesnt exist')
            #print(all_pages)
            # print(all_objects)
            return all_pages


def get_persons(ppn, page):
    all_objects = []
    dir = fulltext_dir + ppn + '\\' + 'FILE_' + page + '_FULLTEXT' + '\\'
    # print(dir)
    json_file_path = dir + page.zfill(8) + '_ner_details.json'
    # print(json_file_path)
    with open(json_file_path) as f:
        data = json.load(f)
        # print(data)
        # print(len(data))
        # print(data['entities'])
        for y in range(0, len(data['entities'])):
            current_element = data['entities'][y]
            # print(current_element)
            current_text = current_element['text']
            if len(current_element['labels']) == 1:
                current_value = current_element['labels'][0]
                print(current_value)
                print(current_text)
                if current_value['_value'] == 'PER':
                    if current_text not in all_objects:
                        all_objects.append(current_text)
            else:
                print('There seems to be more than one label')
    print('all persons = ' + str(all_objects))
    return all_objects


def get_illustrated_persons(ppn):
    persons_dict = {}
    result_get_pages_with_object_illu = get_pages_with_object_illu('person', ppn, True)
    pages_with_persons_illu = result_get_pages_with_object_illu[0]
    used_words = result_get_pages_with_object_illu[1]
    pages_with_persons_text = get_pages_with_named_entity('PER', ppn)
    # print(pages_with_persons_text)
    print(pages_with_persons_illu)
    pages_illu = []
    all_persons = []
    for y in range (0,len(pages_with_persons_illu)):
        persons = []
        current_dir = pages_with_persons_illu[y]
        head, tail = os.path.split(current_dir)
        page = tail[0:4]
        # print(page)
        if page not in pages_illu:
            pages_illu.append(page)
    print(pages_with_persons_text)
    print('pages_illu = ' + str(pages_illu))
    for x in range(0, len(pages_illu)):
        if pages_illu[x] in pages_with_persons_text:
            print("found something")
            # print(pages_illu[x])
            persons = get_persons(ppn, pages_illu[x])
            for b in range(0, len(persons)):
                key = persons[b]
                if key not in persons_dict:
                    persons_dict[key] = []
                    persons_dict[key].append(pages_illu[x])
                else:
                    persons_dict[key].append(pages_illu[x])
        # check doublepage
        else:
            num = int(pages_illu[x])
            if(num % 2) == 0:
                other_page = int(pages_illu[x]) + 1
                other_page = str(other_page)
                other_page = other_page.zfill(4)
                # print(other_page)
            else:
                other_page = int(pages_illu[x]) - 1
                other_page = str(other_page)
                other_page = other_page.zfill(4)
            if other_page in pages_with_persons_text:
                # found some doublepage relation
                persons = get_persons(ppn, other_page)
                for a in range(0, len(persons)):
                    key = persons[a]
                    if key not in persons_dict:
                        persons_dict[key] = []
                        persons_dict[key].append(pages_illu[x])
                    else:
                        persons_dict[key].append(pages_illu[x])
                print('persons = ' + str(persons))
            else:
                print("nothing")
        for z in range (0, len(persons)):
            if persons[z] not in all_persons:
                all_persons.append(persons[z])
    print("dictionary =")
    print(persons_dict)
    print(all_persons)
    return all_persons, persons_dict, used_words


def analyze_persons():
    ppns = []
    with open("text_image\\test-ten-books.txt") as f:
        lines = f.readlines()
        for line in lines:
           ppns.append(line.replace("\n", ""))
        f.close()
    wb = load_workbook('text_image\\wd-parsing.xlsx')
    ws1 = wb["OCR_ger_filtered"]
    ws2 = wb.create_sheet("Illustrated Persons")
    row_counter = 2
    ws2.cell(1, 1).value = "title"
    ws2.cell(1, 2).value = "creator"
    ws2.cell(1, 3).value = "subject"
    ws2.cell(1, 4).value = "publisher"
    ws2.cell(1, 5).value = "type"
    ws2.cell(1, 6).value = "digital object identifier"
    ws2.cell(1, 7).value = "physical object identifier"
    ws2.cell(1, 8).value = "URL"
    ws2.cell(1, 9).value = "language"
    ws2.cell(1, 10).value = "date"
    ws2.cell(1, 11).value = "relation"
    ws2.cell(1, 12).value = "coverage"
    ws2.cell(1, 13).value = "OCR URL"
    ws2.cell(1, 14).value = "OCR found?"
    ws2.cell(1, 15).value = "ddc"
    ws2.cell(1, 16).value = "Illustrated person"
    ws2.cell(1, 17).value = "pages"
    ws2.cell(1, 18).value = "number of pages"
    for cell in ws2[1:1]:
        cell.font = Font(bold=True)
    for x in range(0, len(ppns)):
        illustrated_persons_in_ppn = (get_illustrated_persons(ppns[x]))
        illustrated_persons_in_ppn_list = illustrated_persons_in_ppn[0]
        dictionary = illustrated_persons_in_ppn[1]
        used_classes = illustrated_persons_in_ppn[2]
        used_classes_str = ', '.join(used_classes)
        print(ppns[x])
        print(illustrated_persons_in_ppn)
        for z in range(0, len(illustrated_persons_in_ppn_list)):
            for y in range(1, 3000):
                if ws1.cell(y, 6).value == ppns[x]:
                    ws2.cell(row_counter, 1).value = ws1.cell(y, 1).value
                    ws2.cell(row_counter, 2).value = ws1.cell(y, 2).value
                    ws2.cell(row_counter, 3).value = ws1.cell(y, 3).value
                    ws2.cell(row_counter, 4).value = ws1.cell(y, 4).value
                    ws2.cell(row_counter, 5).value = ws1.cell(y, 5).value
                    ws2.cell(row_counter, 6).value = ws1.cell(y, 6).value
                    ws2.cell(row_counter,  7).value = ws1.cell(y, 7).value
                    ws2.cell(row_counter, 8).value = ws1.cell(y, 8).value
                    ws2.cell(row_counter, 9).value = ws1.cell(y, 9).value
                    ws2.cell(row_counter, 10).value = ws1.cell(y, 10).value
                    ws2.cell(row_counter, 11).value = ws1.cell(y, 11).value
                    ws2.cell(row_counter, 12).value = ws1.cell(y, 12).value
                    ws2.cell(row_counter, 13).value = ws1.cell(y, 13).value
                    ws2.cell(row_counter, 14).value = ws1.cell(y, 14).value
                    ws2.cell(row_counter, 15).value = ws1.cell(y, 15).value
                    ws2.cell(row_counter, 16).value = illustrated_persons_in_ppn_list[z]
                    string_pages = ', '.join(dictionary[illustrated_persons_in_ppn_list[z]])
                    ws2.cell(row_counter, 17).value = string_pages
                    ws2.cell(row_counter, 18).value = len(dictionary[illustrated_persons_in_ppn_list[z]])
                    ws2.cell(row_counter, 19).value = used_classes_str
                    row_counter += 1
    wb.save('text_image\\results.xlsx')


def get_text_length(ppn):
    fulltext_file = 'sbbget\\sbbget_downloads\\download_temp\\' + ppn + '\\' + 'fulltext.txt'
    file = open(fulltext_file, "r", encoding='utf-8')
    number_of_lines = 0
    number_of_words = 0
    number_of_characters = 0
    for line in file:
        line = line.strip("\n")
        words = line.split()
        number_of_lines += 1
        number_of_words += len(words)
        number_of_characters += len(line)
    file.close()
    # print("lines:", number_of_lines, "words:", number_of_words, "characters:", number_of_characters)
    return number_of_words, number_of_characters


def get_number_of_pages(ppn):
    dir = 'sbbget\\sbbget_downloads\download_temp' + '\\' + ppn
    for x in range(1, 9999):
        current_page = str(x).zfill(4)
        current_dir = 'sbbget\\sbbget_downloads\download_temp\\' + ppn + '\FILE_' + current_page + '_FULLTEXT'
        if os.path.isdir(current_dir) == False:
            number_of_pages = x - 1
            # print(ppn)
            # print(number_of_pages)
            return number_of_pages


def get_number_of_illus(ppn):
    number_of_illus = 0
    dir = 'sbbget\\sbbget_downloads\download_temp\\' + ppn
    for x in range(1, 9999):
        current_page = str(x).zfill(4)
        current_page2 = str(x).zfill(8)
        file = dir + '\\' + 'FILE_' + current_page + '_FULLTEXT\\' + current_page2 + '.xml'
        # print(file)
        if os.path.exists(file):
            # print('exists')
            tree = ET.parse(file)
            root = tree.getroot()
            elemList = []
            for elem in tree.iter():
                # print(elem.tag)
                if elem.tag == '{http://www.loc.gov/standards/alto/ns-v2#}Illustration':
                    number_of_illus += 1
        else:
            print(number_of_illus)
            return number_of_illus

# all classes of YOLO appearing in the corpus, + translation in txt file
def img_objects_quantitaitve():
    translator = Translator()
    all_objects = []
    dictionary = {}
    for filename in os.listdir(object_dir):
        f = os.path.join(object_dir, filename)
        with open(f, 'r') as myfile:
            data = myfile.read()
            data = data.replace("\\", "/")
            y = json.loads(data)
            for x in range(0, len(y)):
                element = y[x]
                z = element.get("objects")
                for a in range(0, len(z)):
                    b = z[a]
                    obj = b.get('name')
                    if obj not in all_objects:
                        all_objects.append(obj)
                        print(all_objects)
    a = open("text_image\\all_classes.txt", "w")
    b = open("text_image\\all_classes_ger.txt", "w")
    c = open("text_image\\9k_labels.txt", "r")
    lines = c.readlines()
    for x in range(0, len(all_objects)):
        with open("text_image\\all_classes.txt", "a") as a:
            a.write(all_objects[x] + "\n")
        translated = GoogleTranslator(source='auto', target='de').translate(all_objects[x])
        with open("text_image\\all_classes_ger.txt", "a") as b:
            b.write(translated + "\n")
        print("lines:")
        print(lines)
        synset_id = get_synset_id(all_objects[x])
        a = re.search(r"\d+(\.\d+)?", synset_id)
        id_str = a.group(0)
        id_int = int(id_str)
        synset = wordnet.synset_from_pos_and_offset('n', id_int)
        print(synset.definition())
        with open("text_image\\definitions.txt", "a") as c:
            c.write(synset.definition() + '\n')


def quantitative():
    ppns = []
    with open(ppn_list) as f:
        lines = f.readlines()
        for line in lines:
           ppns.append(line.replace("\n", ""))
        f.close()
    wb = load_workbook('image-analysis\\wd-parsing.xlsx')
    ws1 = wb["OCR_ger"]
    ws2 = wb.create_sheet("Text+Illu quantitativ")
    row_counter = 2
    ws2.cell(1, 1).value = "title"
    ws2.cell(1, 2).value = "creator"
    ws2.cell(1, 3).value = "subject"
    ws2.cell(1, 4).value = "publisher"
    ws2.cell(1, 5).value = "type"
    ws2.cell(1, 6).value = "digital object identifier"
    ws2.cell(1, 7).value = "physical object identifier"
    ws2.cell(1, 8).value = "URL"
    ws2.cell(1, 9).value = "language"
    ws2.cell(1, 10).value = "date"
    ws2.cell(1, 11).value = "relation"
    ws2.cell(1, 12).value = "coverage"
    ws2.cell(1, 13).value = "OCR URL"
    ws2.cell(1, 14).value = "OCR found?"
    ws2.cell(1, 15).value = "number of pages"
    ws2.cell(1, 16).value = "number of words"
    ws2.cell(1, 17).value = "number of characters"
    ws2.cell(1, 18).value = "number of illustrations"
    for cell in ws2[1:1]:
        cell.font = Font(bold=True)
    for x in range (0, len(ppns)):
        number_of_words = get_text_length(ppns[x])[0]
        number_of_characters = get_text_length(ppns[x])[1]
        number_of_pages = get_number_of_pages(ppns[x])
        number_of_illus = get_number_of_illus(ppns[x])
        # print(number_of_words)
        # print(number_of_characters)
        for y in range(1, 3000):
            if ws1.cell(y, 6).value == ppns[x]:
                # print("found row")
                ws2.cell(row_counter, 1).value = ws1.cell(y, 1).value
                ws2.cell(row_counter, 2).value = ws1.cell(y, 2).value
                ws2.cell(row_counter, 3).value = ws1.cell(y, 3).value
                ws2.cell(row_counter, 4).value = ws1.cell(y, 4).value
                ws2.cell(row_counter, 5).value = ws1.cell(y, 5).value
                ws2.cell(row_counter, 6).value = ws1.cell(y, 6).value
                ws2.cell(row_counter, 7).value = ws1.cell(y, 7).value
                ws2.cell(row_counter, 8).value = ws1.cell(y, 8).value
                ws2.cell(row_counter, 9).value = ws1.cell(y, 9).value
                ws2.cell(row_counter, 10).value = ws1.cell(y, 10).value
                ws2.cell(row_counter, 11).value = ws1.cell(y, 11).value
                ws2.cell(row_counter, 12).value = ws1.cell(y, 12).value
                ws2.cell(row_counter, 13).value = ws1.cell(y, 13).value
                ws2.cell(row_counter, 14).value = ws1.cell(y, 14).value
                ws2.cell(row_counter, 15).value = number_of_pages
                ws2.cell(row_counter, 16).value = number_of_words
                ws2.cell(row_counter, 17).value = number_of_characters
                ws2.cell(row_counter, 18).value = number_of_illus
                row_counter += 1
    #wb.save('image-analysis\\results.xlsx')


def get_all_objects():
    ppns = []
    dictionary = {}
    with open(ppn_list) as f:
        lines = f.readlines()
        for line in lines:
            ppns.append(line.replace("\n", ""))
        f.close()
        for x in range(0, len(ppns)):
            objects = get_objects(ppns[x])
            dictionary.update({ppns[x]: objects})
        print(dictionary)


def get_metadata(ppn):
    wb = load_workbook('text_image\\wd-parsing.xlsx')
    ws1 = wb["OCR_ger"]
    for x in range(1,600):
        if ws1.cell(x, 6).value == ppn:
            title = ws1.cell(x, 1).value
            creator = ws1.cell(x, 2).value
            publisher = ws1.cell(x, 4).value
            language = ws1.cell(x, 9).value
            date = ws1.cell(x, 10).value
            coverage = ws1.cell(x, 12).value
            return title, creator, publisher, language, date, coverage


def get_pages_with_object_text(ppn, word_to_search_for, hyponyms):
    pages_with_object = []
    dir_current_ppn = text_dir + '\\' + ppn
    dirs = os.listdir(dir_current_ppn)
    all_hyponyms = get_hyponyms(word_to_search_for, 'text')
    for x in range(0, len(dirs)):
        current_page = dirs[x]
        file_name = dir_current_ppn + '\\' + current_page + '\\' + ppn + '_' + current_page + '_nouns_lemmatized.txt'
        if hyponyms:
            words_to_search_for = all_hyponyms
        else:
            words_to_search_for = word_to_search_for
        with open(file_name, 'r') as f:
            for line in f:
                for word in line.split():
                    if word in words_to_search_for:
                        pages_with_object.append(current_page)
    return pages_with_object

"""
def get_pages_with_object_text_with_synset_id(ppn, id, hyponyms):
    pages_with_object = []
    dir_current_ppn = text_dir + '\\' + ppn
    dirs = os.listdir(dir_current_ppn)
    words_to_search_for = []
    for x in range(0, len(dirs)):
        current_page = dirs[x]
        file_name = dir_current_ppn + '\\' + current_page + '\\' + ppn + '_' + current_page + '_nouns_lemmatized.txt'
        if hyponyms:
            all_hyponyms = get_hyponyms(word_to_search_for, 'text')
            words_to_search_for = all_hyponyms
        else:
            word_to_search_for = word_to_search_for
        with open(file_name, 'r') as f:
            for line in f:
                for word in line.split():
                    if word in words_to_search_for:
                        pages_with_object.append(current_page)
    return pages_with_object
"""

def get_objects_in_page(ppn, page):
    all_objects = []
    file = object_dir + ppn + '_objects.json'
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        y = json.loads(data)
        for x in range(0, len(y)):
            element = y[x]
            filename = element.get("filename")
            head, tail = os.path.split(filename)
            if tail [0:4] == page:
                z = element.get("objects")
                for a in range(0, len(z)):
                    b = z[a]
                    obj = b.get('name')
                    all_objects.append(obj)
    return all_objects


def is_illustrated(ppn, page):
    file = object_dir + ppn + '_objects.json'
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        y = json.loads(data)
        for x in range(0, len(y)):
            element = y[x]
            filename = element.get("filename")
            head, tail = os.path.split(filename)
            if tail[0:4] == page:
                print("this page is illustrated")
                return True
        return False

# print(is_illustrated('PPN740909703', '0002'))

def get_illustrated_pages(ppn):
    all_pages = []
    file = object_dir + ppn + '_objects.json'
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        y = json.loads(data)
        for x in range(0, len(y)):
            element = y[x]
            z = element.get("objects")
            a = element.get("filename")
            head, tail = os.path.split(a)
            page = tail[0:4]
            all_pages.append(page)
    return all_pages


#specific_obj (and hyponyms) + *any* illustration:
def get_obj_plus_illus(object_name_de, hyponyms, doublepage):
    ppns = []
    dict = {}
    with open(ppn_list) as f:
        lines = f.readlines()
        for line in lines:
            ppns.append(line.replace("\n", ""))
        f.close()
    for x in range(0, len(ppns)):
        print(ppns[x])
        illustrated_pages_of_ppn = get_illustrated_pages(ppns[x])
        ppn_metadata = get_metadata(ppns[x])
        # filter for any metadata here
        pages_with_object_text = get_pages_with_object_text(ppns[x], object_name_de, hyponyms)
        print(pages_with_object_text)
        print(illustrated_pages_of_ppn)
        pages_for_ppn = []
        for y in range(0, len(pages_with_object_text)):
            if pages_with_object_text in illustrated_pages_of_ppn:
                pages_for_ppn.append(pages_with_object_text[y])
            else:
                if doublepage:
                # check doublepage
                    num = int(pages_with_object_text[y])
                    if (num % 2) == 0:
                        other_page = int(pages_with_object_text[y]) + 1
                        other_page = str(other_page)
                        other_page = other_page.zfill(4)
                    else:
                        other_page = int(pages_with_object_text[y]) - 1
                        other_page = str(other_page)
                        other_page = other_page.zfill(4)
                    if other_page in illustrated_pages_of_ppn:
                        pages_for_ppn.append(pages_with_object_text[y])
                        print("dict = ")
                        print(dict)
        dict.update({ppns[x]: pages_for_ppn})
        print("dict = ")
        print(dict)
    return dict


#specific_obj (and hyponyms) + sitable illustration:
def get_obj_plus_illus_specific(object_name_de, object_name_en, hyponyms, doublepage):
    ppns = []
    dict = {}
    with open(ppn_list) as f:
        lines = f.readlines()
        for line in lines:
            ppns.append(line.replace("\n", ""))
        f.close()
    for x in range(0, len(ppns)):
        print(ppns[x])
        ppn_metadata = get_metadata(ppns[x])
        # filter for any metadata here
        pages_with_object_text = get_pages_with_object_text(ppns[x], object_name_de, hyponyms)
        pages_with_object_illu = get_pages_with_object_illu((ppns[x]), object_name_en, hyponyms)
        print(pages_with_object_text)
        print(pages_with_object_illu)
        pages_for_ppn = []
        for y in range(0, len(pages_with_object_text)):
            if pages_with_object_text in pages_with_object_illu:
                pages_for_ppn.append(pages_with_object_text[y])
            else:
                if doublepage:
                # check doublepage
                    num = int(pages_with_object_text[y])
                    if (num % 2) == 0:
                        other_page = int(pages_with_object_text[y]) + 1
                        other_page = str(other_page)
                        other_page = other_page.zfill(4)
                    else:
                        other_page = int(pages_with_object_text[y]) - 1
                        other_page = str(other_page)
                        other_page = other_page.zfill(4)
                    if other_page in pages_with_object_illu:
                        pages_for_ppn.append(pages_with_object_text[y])
                        print("dict = ")
                        print(dict)
        dict.update({ppns[x]: pages_for_ppn})
        print("dict = ")
        print(dict)


#get_pages_with_object_text('PPN74082984X', 'Katze', True)
get_obj_plus_illus('Katze', True, True)

# examples:
# get_all_objects()
# get_objects('PPN740909118')
# get_illustrated_persons('PPN741112043')
# analyze_persons()
# quantitative()
# img_objects_quantitaitve()
# number_of_synsets()
# img_objects_quantitaitve()


# experiments with translations ...
"""
en = wn.Wordnet(lexicon='ewn')
de = wn.Wordnet(lexicon='odenet')
synset2 = wn.synsets('Rinderartige', lang='de')
print(synset2)
ger = wn.Wordnet(lexicon='odenet')
synset=wn.synset(id='pwn-00002684-n')
print(synset.lemmas())
translations = synset.translate(lang='de')
print(translations)
for x in range(0, len(translations)):
    print(translations[x].lemmas())
    print(translations[x].id)
    print(translations[x].definition())
    id = translations[x].id
    ili = wn.synset(id).ili
    ss = translations[x]
    hypernyms = ss.hypernyms()
    for x in range(0, len(hypernyms)):
        print(ss.hyponyms()[x].lemmas())
    hypos = ss.hyponyms()
    print(hypos)
    for x in range(0, len(hypos)):
        print(ss.hyponyms()[x].lemmas())
    print(ili)


en = wn.Wordnet(lexicon='ewn')
de = wn.Wordnet(lexicon='odenet')
synset = wn.synset(id='pwn-00002684-n')
print(synset.lemmas())
translation = synset.translate(lang='de')
print(translation)
print(translation[0])
print(translation[0].lemmas())
print(translation[0].hypernyms())
"""