import os
from pathlib import Path
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Font
import re

directory_mets = 'sbbget\sbbget_downloads\download_temp'

wb = load_workbook('analysis/wd-parsing.xlsx')
ws1 = wb['Sheet']
ws1.cell(1, 15).value = "DDC"
ws1.cell(1, 16).value = "DDC2"
ws1.cell(1, 17).value = "extent"
ws1.cell(1, 18).value = "extent"
ws1.cell(1, 19).value = "extent"
for cell in ws1[1:1]:
    cell.font = Font(bold=True)

software_names = []

for subdir, dirs, files in os.walk(directory_mets):
    name_directory = Path(subdir).parts[-1]
    if name_directory == '__metsmods':
        current_ppn = Path(subdir).parts[-2]
        current_mets_file = subdir + '\\' + current_ppn + '.xml'
        tree = ET.parse(current_mets_file)
        root = tree.getroot()
        y = 3000
        for z in range(1, 3000):
            if ws1.cell(z, 6).value == current_ppn:
                y = z
                break
        w = 15
        x = 17
        for elem in tree.iter():
            # print(elem.tag)
            if elem.tag == '{http://www.loc.gov/mods/v3}classification':
                authority_value = elem.attrib.get("authority")
                if authority_value == 'ddc':
                    ddc = elem.text
                    ddc_list = re.findall("\d+\.\d+", ddc)
                    for v in range(0, len(ddc_list)):
                            ws1.cell(y, w).value = float(ddc_list[v])
                            w += 1
    if subdir.endswith('FULLTEXT'):
        current_ppn = Path(subdir).parts[-2]
        list = os.listdir(subdir)
        software_list = []
        for x in range(0, len(list)):
            if list[x].endswith('.xml'):
                path = subdir  + '\\' + list[x]
                tree = ET.parse(path)
                root = tree.getroot()
                software_name = 'unknown'
                software_version = 'unknown'
                for elem in tree.iter():
                    if elem in tree.iter():
                        if elem.tag == '{http://www.loc.gov/standards/alto/ns-v2#}softwareName':
                            software_name = elem.text
                        if elem.tag == '{http://www.loc.gov/standards/alto/ns-v2#}softwareVersion':
                            software_version = elem.text
                software = software_name + ' ' + software_version
                if software not in software_list:
                    software_list.append(software)
                # print(software_list)


wb.save('mets_mods_parsing/wd-parsing.xlsx')







