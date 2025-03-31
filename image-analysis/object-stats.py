import json
import os
from openpyxl import load_workbook
from openpyxl.styles import Font
import xml.etree.ElementTree as ET
from deep_translator import GoogleTranslator
from nltk.corpus import wordnet as wn
from googletrans import Translator
from nltk.corpus import wordnet


object_dir = 'image-analysis\\results\\yolo9000\\'
fulltext_dir = 'sbbget\\sbbget_downloads\download_temp\\'

def all_objects_in_corpus():
    objects_in_corpus = []
    for subdir, dirs, files in os.walk(object_dir):
        for filename in files:
            ppn = filename[0:12]
            objects = get_objects(ppn)
            for x in range(0,len(objects)):
                if objects[x] not in objects_in_corpus:
                    objects_in_corpus.append(objects[x])
                    objects_in_corpus.sort()
                    print(objects_in_corpus)
                    return(objects_in_corpus)


def get_objects(ppn):
    all_objects = []
    file = object_dir + ppn + '_objects.json'
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        y = json.loads(data)
        for x in range(0, len(y)):
            element = y[x]
            z = element.get("objects")
            for a in range(0, len(z)):
                b = z[a]
                obj = b.get('name')
                if obj not in all_objects:
                    all_objects.append(obj)
    print(all_objects)
    return all_objects


def get_pages_with_object_illu(object_name, ppn):
    all_pages = []
    file = object_dir + ppn + '_objects.json'
    # print(file)
    with open(file, 'r') as myfile:
        data = myfile.read()
        data = data.replace("\\", "/")
        # print(data)
        y = json.loads(data)
        # print(y)
        # print(len(y))
        for x in range(0, len(y)):
            element = y[x]
            z = element.get("objects")
            # print(z)
            for a in range(0, len(z)):
                b = z[a]
                # print(b)
                obj = b.get('name')
                # print('obj = ' + obj)
                if obj == object_name:
                    # print("Found something")
                    page = y[x].get('filename')
                    # print(page)
                    all_pages.append(page)
    print(all_pages)
    return(all_pages)


def get_pages_with_object_text(object_name, ppn):
    all_pages = []
    all_objects = []
    current_page = "1"
    for x in range(1, 9999):
        current_page = str(x)
        current_page = current_page.zfill(4)
        dir = fulltext_dir + ppn + '\\' + 'FILE_' + current_page + '_FULLTEXT' + '\\'
        json_file_path = dir + current_page.zfill(8) + '_ner_details.json'
        if os.path.exists(dir):
            if os.path.exists(json_file_path):
                with open(json_file_path) as f:
                    data = json.load(f)
                    for y in range(0, len(data['entities'])):
                        current_element = data['entities'][y]
                        current_text = current_element['text']
                        if len(current_element['labels']) == 1:
                            current_value = current_element['labels'][0]
                            if current_value['_value'] == object_name:
                                if current_page not in all_pages:
                                    all_pages.append(current_page)
                                if current_text not in all_objects:
                                    all_objects.append(current_text)
                        else:
                            print('There seems to be more than one label')
        else:
            return all_pages


def get_persons(ppn, page):
    all_objects = []
    dir = fulltext_dir + ppn + '\\' + 'FILE_' + page + '_FULLTEXT' + '\\'
    # print(dir)
    json_file_path = dir + page.zfill(8) + '_ner_details.json'
    # print(json_file_path)
    with open(json_file_path) as f:
        data = json.load(f)
        # print(data)
        # print(len(data))
        # print(data['entities'])
        for y in range(0, len(data['entities'])):
            current_element = data['entities'][y]
            # print(current_element)
            current_text = current_element['text']
            if len(current_element['labels']) == 1:
                current_value = current_element['labels'][0]
                print(current_value)
                print(current_text)
                if current_value['_value'] == 'PER':
                    if current_text not in all_objects:
                        all_objects.append(current_text)
            else:
                print('There seems to be more than one label')
    print('all persons = ' + str(all_objects))
    return all_objects


def get_illustrated_persons(ppn):
    persons_dict = {}
    pages_with_persons_illu = get_pages_with_object_illu('person', ppn)
    pages_with_persons_text = get_pages_with_object_text('PER', ppn)
    # print(pages_with_persons_text)
    print(pages_with_persons_illu)
    pages_illu = []
    all_persons = []
    for y in range (0,len(pages_with_persons_illu)):
        persons = []
        current_dir = pages_with_persons_illu[y]
        head, tail = os.path.split(current_dir)
        page = tail[0:4]
        # print(page)
        if page not in pages_illu:
            pages_illu.append(page)
    print(pages_with_persons_text)
    print('pages_illu = ' + str(pages_illu))
    for x in range(0, len(pages_illu)):
        if pages_illu[x] in pages_with_persons_text:
            print("found something")
            # print(pages_illu[x])
            persons = get_persons(ppn, pages_illu[x])
            for b in range(0, len(persons)):
                key = persons[b]
                if key not in persons_dict:
                    persons_dict[key] = []
                    persons_dict[key].append(pages_illu[x])
                else:
                    persons_dict[key].append(pages_illu[x])
        # check doublepage
        else:
            num = int(pages_illu[x])
            if(num % 2) == 0:
                other_page = int(pages_illu[x]) + 1
                other_page = str(other_page)
                other_page = other_page.zfill(4)
                # print(other_page)
            else:
                other_page = int(pages_illu[x]) - 1
                other_page = str(other_page)
                other_page = other_page.zfill(4)
            if other_page in pages_with_persons_text:
                # found some doublepage relation
                persons = get_persons(ppn, other_page)
                for a in range(0, len(persons)):
                    key = persons[a]
                    if key not in persons_dict:
                        persons_dict[key] = []
                        persons_dict[key].append(pages_illu[x])
                    else:
                        persons_dict[key].append(pages_illu[x])
                print('persons = ' + str(persons))
            else:
                print("nope")
        for z in range (0, len(persons)):
            if persons[z] not in all_persons:
                all_persons.append(persons[z])
    print("dictionary =")
    print(persons_dict)
    print(all_persons)
    return all_persons, persons_dict


def analyze_persons():
    ppns = []
    with open("image-analysis\\WD_OCR_ger.txt") as f:
        lines = f.readlines()
        for line in lines:
           ppns.append(line.replace("\n", ""))
        f.close()
    wb = load_workbook('image-analysis\\wd-parsing.xlsx')
    ws1 = wb["OCR_ger"]
    ws2 = wb.create_sheet("Illustrated Persons")
    row_counter = 2
    ws2.cell(1, 1).value = "title"
    ws2.cell(1, 2).value = "creator"
    ws2.cell(1, 3).value = "subject"
    ws2.cell(1, 4).value = "publisher"
    ws2.cell(1, 5).value = "type"
    ws2.cell(1, 6).value = "digital object identifier"
    ws2.cell(1, 7).value = "physical object identifier"
    ws2.cell(1, 8).value = "URL"
    ws2.cell(1, 9).value = "language"
    ws2.cell(1, 10).value = "date"
    ws2.cell(1, 11).value = "relation"
    ws2.cell(1, 12).value = "coverage"
    ws2.cell(1, 13).value = "OCR URL"
    ws2.cell(1, 14).value = "OCR found?"
    ws2.cell(1, 15).value = "ddc"
    ws2.cell(1, 16).value = "Illustrated person"
    ws2.cell(1, 17).value = "pages"
    ws2.cell(1, 18).value = "number of pages"
    for cell in ws2[1:1]:
        cell.font = Font(bold=True)
    for x in range(0, len(ppns)):
        illustrated_persons_in_ppn = (get_illustrated_persons(ppns[x]))
        illustrated_persons_in_ppn_list = illustrated_persons_in_ppn[0]
        dictionary = illustrated_persons_in_ppn[1]
        print(ppns[x])
        print(illustrated_persons_in_ppn)
        for z in range(0, len(illustrated_persons_in_ppn_list)):
            for y in range(1, 3000):
                if ws1.cell(y, 6).value == ppns[x]:
                    ws2.cell(row_counter, 1).value = ws1.cell(y, 1).value
                    ws2.cell(row_counter, 2).value = ws1.cell(y, 2).value
                    ws2.cell(row_counter, 3).value = ws1.cell(y, 3).value
                    ws2.cell(row_counter, 4).value = ws1.cell(y, 4).value
                    ws2.cell(row_counter, 5).value = ws1.cell(y, 5).value
                    ws2.cell(row_counter, 6).value = ws1.cell(y, 6).value
                    ws2.cell(row_counter,  7).value = ws1.cell(y, 7).value
                    ws2.cell(row_counter, 8).value = ws1.cell(y, 8).value
                    ws2.cell(row_counter, 9).value = ws1.cell(y, 9).value
                    ws2.cell(row_counter, 10).value = ws1.cell(y, 10).value
                    ws2.cell(row_counter, 11).value = ws1.cell(y, 11).value
                    ws2.cell(row_counter, 12).value = ws1.cell(y, 12).value
                    ws2.cell(row_counter, 13).value = ws1.cell(y, 13).value
                    ws2.cell(row_counter, 14).value = ws1.cell(y, 14).value
                    ws2.cell(row_counter, 15).value = ws1.cell(y, 15).value
                    ws2.cell(row_counter, 16).value = illustrated_persons_in_ppn_list[z]
                    string_pages = ', '.join(dictionary[illustrated_persons_in_ppn_list[z]])
                    ws2.cell(row_counter, 17).value = string_pages
                    ws2.cell(row_counter, 18).value = len(dictionary[illustrated_persons_in_ppn_list[z]])
                    row_counter += 1
    wb.save('image-analysis\\results.xlsx')


def get_text_length(ppn):
    fulltext_file = 'sbbget\\sbbget_downloads\\download_temp\\' + ppn + '\\' + 'fulltext.txt'
    file = open(fulltext_file, "r", encoding='utf-8')
    number_of_lines = 0
    number_of_words = 0
    number_of_characters = 0
    for line in file:
        line = line.strip("\n")
        words = line.split()
        number_of_lines += 1
        number_of_words += len(words)
        number_of_characters += len(line)
    file.close()
    # print("lines:", number_of_lines, "words:", number_of_words, "characters:", number_of_characters)
    return number_of_words, number_of_characters


def get_number_of_pages(ppn):
    for x in range(1, 9999):
        current_page = str(x).zfill(4)
        current_dir = 'sbbget\\sbbget_downloads\download_temp\\' + ppn + '\FILE_' + current_page + '_FULLTEXT'
        if os.path.isdir(current_dir) == False:
            number_of_pages = x - 1
            # print(number_of_pages)
            return number_of_pages


def get_number_of_illus(ppn):
    number_of_illus = 0
    dir = 'sbbget\\sbbget_downloads\download_temp\\' + ppn
    for x in range(1, 9999):
        current_page = str(x).zfill(4)
        current_page2 = str(x).zfill(8)
        file = dir + '\\' + 'FILE_' + current_page + '_FULLTEXT\\' + current_page2 + '.xml'
        # print(file)
        if os.path.exists(file):
            # print('exists')
            tree = ET.parse(file)
            root = tree.getroot()
            elemList = []
            for elem in tree.iter():
                # print(elem.tag)
                if elem.tag == '{http://www.loc.gov/standards/alto/ns-v2#}Illustration':
                    number_of_illus += 1
        else:
            print(number_of_illus)
            return number_of_illus

# all classes of the corpus in a list
def img_objects_quantitaitve():
    translator = Translator()
    all_objects = []
    dictionary = {}
    for filename in os.listdir(object_dir):
        f = os.path.join(object_dir, filename)
        with open(f, 'r') as myfile:
            data = myfile.read()
            data = data.replace("\\", "/")
            y = json.loads(data)
            for x in range(0, len(y)):
                element = y[x]
                z = element.get("objects")
                for a in range(0, len(z)):
                    b = z[a]
                    obj = b.get('name')
                    if obj not in all_objects:
                        all_objects.append(obj)
                        print(all_objects)
    a = open("image-analysis\\all_classes.txt", "w")
    b = open("image-analysis\\all_classes_ger.txt", "w")
    c = open("image-analysis\\all_labels.LABELS", "r")
    lines = c.readlines()
    for x in range(0, len(all_objects)):
        a.write(all_objects[x] + "\n")
    for x in range(0, len(all_objects)):
        translated = GoogleTranslator(source='auto', target='de').translate(all_objects[x])
        wnetID = lines[x]
        print(wnetID)
        b.write(translated + "\n")
    a.close()

#how many synsets?
def number_of_synsets():
    obejects_ger = []
    with open("image-analysis\\all_classes.txt", "r") as f:
        objects = f.readlines()
        for x in range(0, len(objects)):
            print(wordnet.synset_from_pos_and_offset('n',13354420))
            translator = Translator()
            result = translator.translate(objects[x], src='en', dest='de')
            print(result.text)
            obejects_ger.append(result)


def quantitative():
    ppns = []
    with open("image-analysis/WD_OCR_ger.txt") as f:
        lines = f.readlines()
        for line in lines:
           ppns.append(line.replace("\n", ""))
        f.close()
    wb = load_workbook('image-analysis\\wd-parsing.xlsx')
    ws1 = wb["OCR_ger"]
    ws2 = wb.create_sheet("Text+Illu quantitativ")
    row_counter = 2
    ws2.cell(1, 1).value = "title"
    ws2.cell(1, 2).value = "creator"
    ws2.cell(1, 3).value = "subject"
    ws2.cell(1, 4).value = "publisher"
    ws2.cell(1, 5).value = "type"
    ws2.cell(1, 6).value = "digital object identifier"
    ws2.cell(1, 7).value = "physical object identifier"
    ws2.cell(1, 8).value = "URL"
    ws2.cell(1, 9).value = "language"
    ws2.cell(1, 10).value = "date"
    ws2.cell(1, 11).value = "relation"
    ws2.cell(1, 12).value = "coverage"
    ws2.cell(1, 13).value = "OCR URL"
    ws2.cell(1, 14).value = "OCR found?"
    ws2.cell(1, 15).value = "number of pages"
    ws2.cell(1, 16).value = "number of words"
    ws2.cell(1, 17).value = "number of characters"
    ws2.cell(1, 18).value = "number of illustrations"
    for cell in ws2[1:1]:
        cell.font = Font(bold=True)
    for x in range (0, len(ppns)):
        number_of_words = get_text_length(ppns[x])[0]
        number_of_characters = get_text_length(ppns[x])[1]
        number_of_pages = get_number_of_pages(ppns[x])
        number_of_illus = get_number_of_illus(ppns[x])
        # print(number_of_words)
        # print(number_of_characters)
        for y in range(1, 3000):
            if ws1.cell(y, 6).value == ppns[x]:
                # print("found row")
                ws2.cell(row_counter, 1).value = ws1.cell(y, 1).value
                ws2.cell(row_counter, 2).value = ws1.cell(y, 2).value
                ws2.cell(row_counter, 3).value = ws1.cell(y, 3).value
                ws2.cell(row_counter, 4).value = ws1.cell(y, 4).value
                ws2.cell(row_counter, 5).value = ws1.cell(y, 5).value
                ws2.cell(row_counter, 6).value = ws1.cell(y, 6).value
                ws2.cell(row_counter, 7).value = ws1.cell(y, 7).value
                ws2.cell(row_counter, 8).value = ws1.cell(y, 8).value
                ws2.cell(row_counter, 9).value = ws1.cell(y, 9).value
                ws2.cell(row_counter, 10).value = ws1.cell(y, 10).value
                ws2.cell(row_counter, 11).value = ws1.cell(y, 11).value
                ws2.cell(row_counter, 12).value = ws1.cell(y, 12).value
                ws2.cell(row_counter, 13).value = ws1.cell(y, 13).value
                ws2.cell(row_counter, 14).value = ws1.cell(y, 14).value
                ws2.cell(row_counter, 15).value = number_of_pages
                ws2.cell(row_counter, 16).value = number_of_words
                ws2.cell(row_counter, 17).value = number_of_characters
                ws2.cell(row_counter, 18).value = number_of_illus
                row_counter += 1
    #wb.save('image-analysis\\results.xlsx')



#get_objects('PPN740909118')
#get_pages_with_object_illu('person', 'PPN740909118')
# get_pages_with_object_text('LOC', 'PPN741220075')
# get_illustrated_persons('PPN741112043')
#analyze_persons()
# quantitative()
# img_objects_quantitaitve()
#number_of_synsets()
#img_objects_quantitaitve()